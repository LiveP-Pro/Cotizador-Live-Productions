#!/usr/bin/env python3
"""Generate a literal service catalog from the user's current workbook only.

Every populated service cell belongs to exactly one presentation row. Verification
reads the workbook again and compares both presentation and source trace before
writing an output. No inventory or prior catalog is used as an input.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


DEFAULT_VERSION = "20261014-dj-audio-01"
BOOK_KEY = "servicios"
STRUCTURAL_LABELS = {"cuadro de montaje", "equipo a llevar live productions"}


def populated(value: Any) -> bool:
    # Zero and whitespace are source information, not empty cells.
    return value is not None and value != ""


def literal_text(value: Any) -> str:
    return "" if value is None else str(value)


def display_text(value: Any) -> str:
    return re.sub(r"\s+", " ", literal_text(value)).strip()


def comparable_text(value: Any) -> str:
    text = unicodedata.normalize("NFD", display_text(value))
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    text = re.sub(r"[^\w\s]", " ", text.lower())
    return re.sub(r"\s+", " ", text).strip()


def slug(value: Any) -> str:
    # Keep the existing service-id algorithm; display strings stay literal.
    text = unicodedata.normalize("NFD", display_text(value))
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"[^a-zA-Z0-9]+", "-", text.lower()).strip("-") or "servicio"


def json_value(value: Any) -> Any:
    return int(value) if isinstance(value, float) and value.is_integer() else value


def is_quantity(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def is_black_fill(cell: Any) -> bool:
    if cell.fill.fill_type != "solid":
        return False
    color = cell.fill.fgColor
    if color.type == "rgb":
        return str(color.rgb).upper() in {"000000", "FF000000"}
    return color.type == "indexed" and color.indexed in {0, 8}


def parent_heading(ws: Any, row: int, start_col: int, heading: str) -> bool:
    merged = any(m.min_row <= row <= m.max_row and m.min_col <= start_col
                 and m.max_col >= start_col + 1 for m in ws.merged_cells.ranges)
    letters = "".join(char for char in heading if char.isalpha())
    return merged or bool(letters and letters == letters.upper())


def service_starts(ws: Any) -> list[int]:
    # CEREMONIA has no service-title cell; B1 is a structural header.
    return [1] if ws.title == "CEREMONIA" else [c.column for c in ws[1] if populated(c.value)]


def cells_for_row(ws: Any, row: int, start_col: int, end_col: int) -> list[dict[str, Any]]:
    return [{"address": c.coordinate, "ref": f"{ws.title}!{c.coordinate}", "value": json_value(c.value)}
            for c in next(ws.iter_rows(min_row=row, max_row=row, min_col=start_col, max_col=end_col))
            if not isinstance(c, MergedCell) and populated(c.value)]


def row_type(ws: Any, row: int, start_col: int, name: str) -> str:
    q_cell, e_cell = ws.cell(row, start_col), ws.cell(row, start_col + 1)
    quantity, equipment = q_cell.value, e_cell.value
    if is_quantity(quantity) and populated(equipment):
        return "item"
    if row == 1 and ws.title != "CEREMONIA" and quantity == name:
        return "header"
    heading = equipment if populated(equipment) else (quantity if not is_quantity(quantity) else None)
    if populated(heading) and comparable_text(heading) in STRUCTURAL_LABELS:
        return "header"
    if populated(heading) and (is_black_fill(q_cell) or is_black_fill(e_cell)):
        return "category"
    if is_quantity(quantity) and not populated(equipment):
        return "unlabeledQuantity"
    return "note"


def section_id(version: str, sheet: str, column: int, row: int) -> str:
    return f"{version}:{slug(sheet)}:{get_column_letter(column)}{row}"


def parse_service(ws: Any, start_col: int, end_col: int, name: str,
                  import_source: dict[str, Any]) -> dict[str, Any]:
    last_row = max((c.row for row in ws.iter_rows(min_col=start_col, max_col=end_col)
                    for c in row if populated(c.value)), default=1)
    sections: list[dict[str, Any]] = []
    trace: list[dict[str, Any]] = []
    parent = child = ""

    def start_section(row: int, title: str) -> dict[str, Any]:
        section = {"id": section_id(import_source["version"], ws.title, start_col, row),
                   "title": title, "sourceStartRow": row,
                   "source": f"{ws.title}!{get_column_letter(start_col)}{row}",
                   "importSource": dict(import_source), "items": [], "rows": []}
        sections.append(section)
        return section

    section = start_section(1, name)
    for row in range(1, last_row + 1):
        cells = cells_for_row(ws, row, start_col, end_col)
        if not cells:
            continue
        trace.extend(cells)
        kind = row_type(ws, row, start_col, name)
        quantity, equipment = ws.cell(row, start_col).value, ws.cell(row, start_col + 1).value
        if kind == "category":
            heading = literal_text(equipment if populated(equipment) else quantity)
            if parent_heading(ws, row, start_col, heading):
                parent, child = heading, ""
            else:
                child = heading
            section = start_section(row, " / ".join(part for part in (parent, child) if part))
        rendered = {"type": kind, "sourceType": "data" if kind == "item" else kind,
                    "sourceRow": row, "cells": cells,
                    "description": " · ".join(literal_text(c["value"]) for c in cells)}
        if kind == "item":
            description = literal_text(equipment)
            rendered.update({"sourceItemIndex": len(section["items"]), "quantity": json_value(quantity),
                             "equipment": description, "description": description,
                             "quantityCell": ws.cell(row, start_col).coordinate,
                             "equipmentCell": ws.cell(row, start_col + 1).coordinate})
            notes = [literal_text(c["value"]) for c in cells if ws[c["address"]].column > start_col + 1]
            if notes:
                rendered["notes"] = notes
            section["items"].append([json_value(quantity), description])
        elif kind == "unlabeledQuantity":
            rendered.update({"quantity": json_value(quantity), "quantityCell": ws.cell(row, start_col).coordinate})
        section["rows"].append(rendered)
    return {"name": name, "source": f"{ws.title}!{get_column_letter(start_col)}:{get_column_letter(end_col)}",
            "sourceCells": trace, "importSource": dict(import_source), "mainSections": sections, "extras": []}


def build_catalog(workbook: Path, version: str = DEFAULT_VERSION,
                  imported_at: str | None = None) -> dict[str, Any]:
    wb = openpyxl.load_workbook(workbook, data_only=False)
    sheets = [name for name in wb.sheetnames if name != "INVENTARIO"]
    if not sheets:
        raise ValueError("El archivo no contiene hojas de servicios")
    import_source = {"fileName": workbook.name, "bookKey": BOOK_KEY,
                     "sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
                     "version": version, "authoritative": True}
    imported_at = imported_at or datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    services: dict[str, dict[str, Any]] = {}
    groups: list[dict[str, Any]] = []
    used_ids: Counter[str] = Counter()
    for sheet_name in sheets:
        ws = wb[sheet_name]
        starts = service_starts(ws)
        if not starts or starts[0] != 1:
            raise ValueError(f"Cuadros no identificables desde la primera columna: {sheet_name}")
        last_col = max((c.column for row in ws for c in row if populated(c.value)), default=2)
        service_ids = []
        for index, start_col in enumerate(starts):
            end_col = starts[index + 1] - 1 if index + 1 < len(starts) else max(start_col + 1, last_col)
            name = literal_text(ws.cell(1, start_col).value) if sheet_name != "CEREMONIA" else sheet_name
            base_id = slug(name)
            used_ids[base_id] += 1
            service_id = base_id if used_ids[base_id] == 1 else f"{base_id}-{used_ids[base_id]}"
            services[service_id] = parse_service(ws, start_col, end_col, name, import_source)
            service_ids.append(service_id)
        groups.append({"label": sheet_name, "serviceIds": service_ids})
    return {"version": version,
            "source": {**import_source, "userFile": workbook.name, "workbook": workbook.name,
                       "importedAt": imported_at, "sheets": sheets, "serviceSheets": sheets},
            "services": services, "groups": groups}


def verify_catalog(catalog: dict[str, Any], workbook: Path) -> dict[str, Any]:
    """Compare exact workbook values with trace AND presentation independently.

    Counts alone cannot detect substitutions or a trace that is never rendered.
    Require per-row cell order, literal item pairs, source category boundaries,
    section-local item indices and unique references before publishing an output.
    """
    wb = openpyxl.load_workbook(workbook, data_only=False)
    sheets = [name for name in wb.sheetnames if name != "INVENTARIO"]
    metadata = {"fileName": workbook.name, "bookKey": BOOK_KEY,
                "sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
                "version": catalog["version"], "authoritative": True}
    for key, value in metadata.items():
        if catalog["source"].get(key) != value:
            raise ValueError(f"Procedencia incorrecta: {key}")
    if catalog["source"].get("sheets") != sheets or catalog["source"].get("serviceSheets") != sheets:
        raise ValueError("Hojas distintas de los servicios del archivo actual")

    expected_cells: dict[str, Any] = {}
    traced: dict[str, Any] = {}
    presented: dict[str, Any] = {}
    duplicate_refs: list[str] = []
    types: Counter[str] = Counter()
    items = zero_items = zero_cells = section_count = 0
    reports: list[dict[str, Any]] = []
    expected_groups: list[dict[str, Any]] = []
    used_ids: Counter[str] = Counter()
    expected_ids: list[str] = []
    for sheet_name in sheets:
        ws = wb[sheet_name]
        # Scan the full worksheet rather than reusing generator block/trace data.
        full_cells = [c for row in ws for c in row if c.value is not None and c.value != ""]
        for c in full_cells:
            expected_cells[f"{sheet_name}!{c.coordinate}"] = json_value(c.value)
            zero_cells += int(is_quantity(c.value) and c.value == 0)
        starts = [1] if sheet_name == "CEREMONIA" else [c.column for c in ws[1] if c.value is not None and c.value != ""]
        last_col = max((c.column for c in full_cells), default=2)
        sheet_ids = []
        for index, start_col in enumerate(starts):
            end_col = starts[index + 1] - 1 if index + 1 < len(starts) else max(start_col + 1, last_col)
            source = f"{sheet_name}!{get_column_letter(start_col)}:{get_column_letter(end_col)}"
            name = literal_text(ws.cell(1, start_col).value) if sheet_name != "CEREMONIA" else sheet_name
            base_id = slug(name)
            used_ids[base_id] += 1
            service_id = base_id if used_ids[base_id] == 1 else f"{base_id}-{used_ids[base_id]}"
            expected_ids.append(service_id)
            sheet_ids.append(service_id)
            service = catalog["services"].get(service_id)
            if not service or service.get("source") != source or service.get("name") != name:
                raise ValueError(f"Cuadro ausente o alterado: {source}")
            if service.get("importSource") != metadata:
                raise ValueError(f"Procedencia del cuadro alterada: {source}")
            for c in service.get("sourceCells", []):
                ref = f"{sheet_name}!{c['address']}"
                if c.get("ref") != ref or ref in traced:
                    duplicate_refs.append(ref)
                traced[ref] = c["value"]

            block_cells = [c for c in full_cells if start_col <= c.column <= end_col]
            expected_rows = sorted({c.row for c in block_cells})
            actual_rows = [(s, r) for s in service.get("mainSections", []) for r in s.get("rows", [])]
            if [r.get("sourceRow") for _, r in actual_rows] != expected_rows:
                raise ValueError(f"Filas omitidas, repetidas o fuera de orden: {source}")
            expected_sections: list[tuple[int, str]] = [(1, name)]
            parent = child = ""
            expected_items = []
            actual_items = [item for s in service["mainSections"] for item in s.get("items", [])]
            block_zero = block_categories = 0
            for section, row in actual_rows:
                source_row = row["sourceRow"]
                row_cells = [c for c in block_cells if c.row == source_row]
                row_values = [{"address": c.coordinate, "ref": f"{sheet_name}!{c.coordinate}",
                               "value": json_value(c.value)} for c in row_cells]
                if row.get("cells") != row_values:
                    raise ValueError(f"Celdas visibles alteradas: {sheet_name}, fila {source_row}, cuadro {start_col}")
                for c in row["cells"]:
                    if c["ref"] in presented:
                        duplicate_refs.append(c["ref"])
                    presented[c["ref"]] = c["value"]
                q_cell, e_cell = ws.cell(source_row, start_col), ws.cell(source_row, start_col + 1)
                quantity, equipment = q_cell.value, e_cell.value
                has_equipment = equipment is not None and equipment != ""
                numeric = isinstance(quantity, (int, float)) and not isinstance(quantity, bool)
                heading = equipment if has_equipment else (quantity if not numeric else None)
                if numeric and has_equipment:
                    kind = "item"
                elif (source_row == 1 and sheet_name != "CEREMONIA" and quantity == name) or (
                    heading is not None and comparable_text(heading) in STRUCTURAL_LABELS):
                    kind = "header"
                elif heading is not None and heading != "" and (is_black_fill(q_cell) or is_black_fill(e_cell)):
                    kind = "category"
                elif numeric and not has_equipment:
                    kind = "unlabeledQuantity"
                else:
                    kind = "note"
                if row.get("type") != kind or row.get("sourceType") != ("data" if kind == "item" else kind):
                    raise ValueError(f"Tipo de fila incorrecto: {sheet_name}, fila {source_row}, cuadro {start_col}")
                types[kind] += 1
                if kind == "item":
                    item = [json_value(quantity), literal_text(equipment)]
                    expected_items.append(item)
                    item_index = row.get("sourceItemIndex")
                    if not isinstance(item_index, int) or item_index < 0 or item_index >= len(section["items"]):
                        raise ValueError(f"Indice de equipo invalido: {sheet_name}!{e_cell.coordinate}")
                    if section["items"][item_index] != item or row.get("quantity") != item[0] or row.get("description") != item[1] or row.get("equipment") != item[1]:
                        raise ValueError(f"Equipo o cantidad alterado: {sheet_name}!{e_cell.coordinate}")
                    if row.get("quantityCell") != q_cell.coordinate or row.get("equipmentCell") != e_cell.coordinate:
                        raise ValueError(f"Referencia de equipo incorrecta: {sheet_name}!{e_cell.coordinate}")
                    if row.get("notes", []) != [literal_text(c.value) for c in row_cells if c.column > start_col + 1]:
                        raise ValueError(f"Notas de equipo alteradas: {sheet_name}!{e_cell.coordinate}")
                    block_zero += int(quantity == 0)
                else:
                    if row.get("description") != " · ".join(literal_text(c["value"]) for c in row_values):
                        raise ValueError(f"Texto visible incompleto: {sheet_name}, fila {source_row}, cuadro {start_col}")
                    if kind == "unlabeledQuantity" and row.get("quantity") != json_value(quantity):
                        raise ValueError(f"Cantidad sin descripcion alterada: {sheet_name}!{q_cell.coordinate}")
                    if kind == "category":
                        block_categories += 1
                        heading_text = literal_text(heading)
                        letters = "".join(c for c in heading_text if c.isalpha())
                        merged = any(m.min_row <= source_row <= m.max_row and m.min_col <= start_col
                                     and m.max_col >= start_col + 1 for m in ws.merged_cells.ranges)
                        if merged or (letters and letters == letters.upper()):
                            parent, child = heading_text, ""
                        else:
                            child = heading_text
                        expected_sections.append((source_row, " / ".join(v for v in (parent, child) if v)))
                if section.get("sourceStartRow") != expected_sections[-1][0]:
                    raise ValueError(f"Categoria de fila alterada: {sheet_name}, fila {source_row}, cuadro {start_col}")
            if actual_items != expected_items:
                raise ValueError(f"Equipos omitidos, combinados o fuera de orden: {source}")
            sections = service["mainSections"]
            if [(s.get("sourceStartRow"), s.get("title")) for s in sections] != expected_sections:
                raise ValueError(f"Categorias omitidas o alteradas: {source}")
            for section in sections:
                if section.get("id") != section_id(catalog["version"], sheet_name, start_col, section["sourceStartRow"]):
                    raise ValueError(f"Identidad de seccion incorrecta: {source}")
                if section.get("importSource") != metadata:
                    raise ValueError(f"Procedencia de seccion incorrecta: {source}")
                if [r["sourceItemIndex"] for r in section["rows"] if r["type"] == "item"] != list(range(len(section["items"]))):
                    raise ValueError(f"Indices de equipo repetidos o fuera de orden: {section['id']}")
                if any(k in section for k in ("notes", "sourceNotes", "blankQuantityItems", "unlabeledQuantities")):
                    raise ValueError(f"Informacion visible duplicada: {section['id']}")
            items += len(expected_items)
            zero_items += block_zero
            section_count += len(sections)
            reports.append({"id": service_id, "sheet": sheet_name, "name": name, "source": source,
                            "sourceCells": len(block_cells), "items": len(expected_items),
                            "zeroQuantityItems": block_zero, "categoryRows": block_categories})
        expected_groups.append({"label": sheet_name, "serviceIds": sheet_ids})
    if list(catalog["services"]) != expected_ids or catalog.get("groups") != expected_groups:
        raise ValueError("Cuadros u hojas omitidos, adicionales o fuera de orden")
    if duplicate_refs:
        raise ValueError(f"Referencias repetidas: {', '.join(duplicate_refs[:20])}")
    for label, actual in (("traza", traced), ("presentacion", presented)):
        missing = sorted(expected_cells.keys() - actual.keys())
        extra = sorted(actual.keys() - expected_cells.keys())
        changed = sorted(r for r in expected_cells.keys() & actual.keys() if expected_cells[r] != actual[r])
        if missing or extra or changed:
            raise ValueError(f"Comparacion celda a celda ({label}): faltan={missing[:20]}, adicionales={extra[:20]}, alteradas={changed[:20]}")
    return {"passed": True, "version": catalog["version"], "source": dict(catalog["source"]),
            "serviceSheets": len(sheets), "services": len(expected_ids), "sourceCells": len(expected_cells),
            "zeroSourceCells": zero_cells, "items": items, "zeroQuantityItems": zero_items,
            "sections": section_count, "rowsByType": dict(types), "missingCells": [], "changedCells": [],
            "duplicateReferences": [],
            "checks": {"sourceCellValues": True, "presentationCellValues": True, "rowOrder": True,
                       "literalEquipmentQuantities": True, "sourceCategories": True,
                       "noDuplicateReferences": True, "serviceOnly": True}, "tables": reports}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--version", default=DEFAULT_VERSION)
    parser.add_argument("--imported-at", help="Current upload/import timestamp; default: current UTC time")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    catalog = build_catalog(args.workbook, args.version, args.imported_at)
    report = verify_catalog(catalog, args.workbook)
    # Verify before replacing any output.
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text("window.requerimientoEquipoCatalog = " + json.dumps(catalog, ensure_ascii=False, indent=2) + ";\n", encoding="utf-8")
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: report[k] for k in ("passed", "serviceSheets", "services", "sourceCells", "zeroSourceCells", "items", "zeroQuantityItems", "sections", "rowsByType")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
