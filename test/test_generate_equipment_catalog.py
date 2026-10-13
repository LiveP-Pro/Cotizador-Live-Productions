"""Source-preservation and fail-closed checks for workbook generation.

Run: python -m unittest discover -s test -p 'test_generate_equipment_catalog.py'
"""

import copy
import importlib.util
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill


SPEC = importlib.util.spec_from_file_location(
    "generate_equipment_catalog", Path(__file__).resolve().parents[1] / "scripts/generate-equipment-catalog.py"
)
GENERATOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(GENERATOR)


class LiteralWorkbookGenerationTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.workbook = Path(self.directory.name) / "current-user-workbook.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SERVICIOS"
        ws["A1"], ws["D1"] = "Servicio  literal ", "Otro servicio"
        ws["B4"], ws["A5"], ws["B5"] = "CUADRO DE MONTAJE", "Cantidad", "Consola"
        ws["A6"], ws["B6"] = 0, " Nombre literal \ncon espacios "
        ws["B7"], ws["B8"] = "MICROFONIA", "Voces"
        ws["A9"], ws["B9"] = 1, "Equipo repetido"
        ws["A10"], ws["B10"] = 2, "Equipo repetido"
        ws["D3"], ws["F11"] = 0, " Nota de fuente  literal "
        black = PatternFill(fill_type="solid", fgColor="FF000000")
        for address in ["B4", "A5", "B5", "B7", "B8"]:
            ws[address].fill = black
        ceremony = wb.create_sheet("CEREMONIA")
        ceremony["B1"] = "CUADRO DE MONTAJE"
        ceremony["B1"].fill = black
        ceremony["A2"], ceremony["B2"] = 1, "Equipo ceremonial"
        # A workbook without inventory must generate and verify normally.
        wb.save(self.workbook)
        self.catalog = GENERATOR.build_catalog(self.workbook, imported_at="2026-10-13T00:00:00Z")

    def rows(self, catalog=None):
        return [row for service in (catalog or self.catalog)["services"].values()
                for section in service["mainSections"] for row in section["rows"]]

    def test_literal_values_zero_rows_duplicate_equipment_and_ceremony(self):
        report = GENERATOR.verify_catalog(self.catalog, self.workbook)
        self.assertEqual((report["services"], report["items"], report["zeroSourceCells"], report["zeroQuantityItems"]), (3, 4, 2, 1))
        self.assertEqual(report["rowsByType"]["unlabeledQuantity"], 1)
        first = self.catalog["services"]["servicio-literal"]
        self.assertEqual(first["name"], "Servicio  literal ")
        items = [item for section in first["mainSections"] for item in section["items"]]
        self.assertEqual(items, [[0, " Nombre literal \ncon espacios "], [1, "Equipo repetido"], [2, "Equipo repetido"]])
        note = next(row for row in self.rows() if row["type"] == "note")
        self.assertEqual(note["description"], " Nota de fuente  literal ")
        self.assertEqual(note["cells"][0]["ref"], "SERVICIOS!F11")
        self.assertNotIn("sourceFileUrl", first["importSource"])

    def test_missing_visible_row_fails_even_when_source_trace_is_complete(self):
        changed = copy.deepcopy(self.catalog)
        section = next(s for s in changed["services"]["servicio-literal"]["mainSections"] if s["items"])
        section["rows"] = [row for row in section["rows"] if row["type"] != "item"]
        with self.assertRaisesRegex(ValueError, "Filas omitidas"):
            GENERATOR.verify_catalog(changed, self.workbook)

    def test_quantity_or_literal_description_substitution_fails(self):
        for key, value in [("quantity", 99), ("description", "Normalized name")]:
            with self.subTest(key=key):
                changed = copy.deepcopy(self.catalog)
                next(row for row in self.rows(changed) if row["type"] == "item")[key] = value
                with self.assertRaisesRegex(ValueError, "Equipo o cantidad alterado"):
                    GENERATOR.verify_catalog(changed, self.workbook)

    def test_duplicate_reference_fails(self):
        changed = copy.deepcopy(self.catalog)
        trace = changed["services"]["servicio-literal"]["sourceCells"]
        trace.append(copy.deepcopy(trace[0]))
        with self.assertRaisesRegex(ValueError, "Referencias repetidas"):
            GENERATOR.verify_catalog(changed, self.workbook)

    def test_category_substitution_fails(self):
        changed = copy.deepcopy(self.catalog)
        section = next(s for s in changed["services"]["servicio-literal"]["mainSections"] if s["title"] == "MICROFONIA / Voces")
        section["title"] = "Unrelated category"
        with self.assertRaisesRegex(ValueError, "Categorias omitidas o alteradas"):
            GENERATOR.verify_catalog(changed, self.workbook)

    def test_changed_source_cell_fails(self):
        changed = copy.deepcopy(self.catalog)
        trace = changed["services"]["servicio-literal"]["sourceCells"]
        trace[0] = {**trace[0], "value": "Wrong service"}
        with self.assertRaisesRegex(ValueError, "Comparacion celda a celda"):
            GENERATOR.verify_catalog(changed, self.workbook)


if __name__ == "__main__":
    unittest.main()
